# -*- coding: utf-8 -*-
"""Build the SIH 'Technical Approach' slide as an editable .pptx.

Laid out to the reference the author supplied: five numbered blocks down
the left, the architecture on the right, a summary band across the foot.

Everything drawn on the right is something the prototype actually does.
What the author's own flow diagram showed but the build does not yet have
- DigiLocker verification, Treasury/PFMS release, expert-panel scoring, a
sanction builder - is set apart in its own strip and labelled as planned,
because a judge can open the site and check.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "/home/lenin/Apps Developed/SIH 26136/slide-assets/technical-approach.pptx"

INK    = RGBColor(0x1F, 0x1F, 0x1F)
MUTED  = RGBColor(0x44, 0x44, 0x44)
ACCENT = RGBColor(0xC5, 0x5A, 0x11)          # the reference's highlight
LINE   = RGBColor(0x59, 0x59, 0x59)

PLATFORM = RGBColor(0xBD, 0xD7, 0xEE)        # platform module
EXTERNAL = RGBColor(0xD9, 0xD9, 0xD9)        # external system
LEGAL    = RGBColor(0xFF, 0xE6, 0x99)        # legal instrument
CHECK    = RGBColor(0xE4, 0xD7, 0xF5)        # governance checkpoint
PLANNED  = RGBColor(0xF2, 0xF2, 0xF2)        # not built yet

BLOCKS = [
 ("1. Sealed challenge authoring", RGBColor(0xE7,0xF0,0xFA),
  [("The department writes an ", 0), ("outcome and its success KPIs", 1),
   (", never a product. The criteria set is canonicalised and hashed with ", 0),
   ("SHA-256", 1), (", then published before any solution is seen — so the "
   "target cannot move afterwards.", 0)]),
 ("2. Risk-tiered eligibility", RGBColor(0xE6,0xF4,0xEA),
  [("Five axes — blast radius, data exposure, reversibility, spend and "
    "dependency — compute a ", 0), ("relaxation envelope", 1),
   (". Firms are screened on the ", 0), ("capped risk rather than turnover", 1),
   (", which is the condition ", 0), ("Rule 173(i)", 1), (" attaches.", 0)]),
 ("3. Two contracts, one gate", RGBColor(0xFD,0xEC,0xEC),
  [("An ", 0), ("Evidence Contract", 1), (" (MSInS, milestone-paid) runs the pilot "
    "in a ", 0), ("masked, DPDP-compliant sandbox", 1),
   (". A validation gate separates it from the ", 0), ("Deployment Contract", 1),
   (", so the department only ever buys a specified product.", 0)]),
 ("4. Independent validation", RGBColor(0xF3,0xEA,0xF9),
  [("The validator is ", 0), ("named at step 1, before any result existed", 1),
   (". They ", 0), ("recompute the seal", 1), (" over the published criteria and "
    "check the evidence against it. Edit one criterion and the ", 0),
   ("seal visibly breaks", 1), (".", 0)]),
 ("5. Lawful route router", RGBColor(0xFD,0xF6,0xE3),
  [("From facts already held the platform picks ", 0),
   ("Tier 1 (PAC under 166(i)), Tier 2 (limited tender) or Tier 3 (GeM "
    "catalogue)", 1), (" — or states plainly that ", 0),
   ("no lawful route exists", 1), (". It will not manufacture one.", 0)]),
]

STAGES = [
 ("1  CHALLENGE AUTHORING",
  [("Department\nproblem", EXTERNAL), ("Challenge Studio\noutcome + success KPIs", PLATFORM),
   ("Clause Injector\nGFR 173(i)", PLATFORM), ("KPI Seal\nSHA-256 + timestamp", PLATFORM),
   ("Published challenge\npathway declared", LEGAL)]),
 ("2  DISCOVERY AND SELECTION",
  [("Startup registries\nStartup India / MSInS", EXTERNAL),
   ("Screening Engine\nrisk-tiered eligibility", PLATFORM),
   ("Eligible?", None), ("Shortlist\ntaken into the pilot", CHECK)]),
 ("3  PILOT EXECUTION",
  [("Evidence Contract\nMSInS, milestone-based", LEGAL),
   ("Sandbox Provisioner\nmasked data, DPDP", PLATFORM),
   ("KPI Tracker\nmilestone evidence", PLATFORM),
   ("Milestone released\nevidence on file", EXTERNAL)]),
 ("4  VALIDATION AND SCALE-UP",
  [("Independent validation\npre-named reviewer", CHECK), ("KPIs met?", None),
   ("Tier Router\n166(i) / limited / GeM", PLATFORM),
   ("Deployment Contract\nthe department buys", LEGAL)]),
]

PLANNED_ITEMS = ["DigiLocker entity verification", "Expert-panel scoring",
                 "Sanction builder / audit packet", "Treasury–PFMS release"]


def textbox(sl, x, y, w, h, runs, size=10.5, bold_first=False, align=PP_ALIGN.LEFT,
            colour=INK, space=1.08):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = space
    for txt, emph in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.name = "Calibri"
        r.font.bold = bool(emph)
        r.font.color.rgb = ACCENT if emph else colour
    return tb


def node(sl, x, y, w, h, label, fill):
    if fill is None:                                   # a decision
        sh = sl.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    else:
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.adjustments[0] = 0.12
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = label.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER; p.line_spacing = 0.92
        r = p.add_run(); r.text = ln
        small = (fill is None)                     # wrap a one-word decision
        r.font.name = "Calibri"
        r.font.size = Pt((7.8 if small else 8.5) if i == 0 else 7.2)
        r.font.bold = (i == 0); r.font.color.rgb = INK
    return sh


def arrow(sl, x1, y, x2):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y),
                                Inches(x2), Inches(y))
    c.line.color.rgb = LINE; c.line.width = Pt(1.25)
    el = c.line._get_or_add_ln()
    el.append(__import__("pptx").oxml.parse_xml(
        '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' type="triangle" w="med" len="med"/>'))
    return c


def frame(sl, x, y, w, h, title):
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    sh.fill.background()
    sh.line.color.rgb = RGBColor(0x9A, 0x9A, 0x9A)
    sh.line.width = Pt(0.75); sh.line.dash_style = 4        # dashed
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    textbox(sl, x + 0.10, y + 0.04, 4.0, 0.2,
            [(title, 0)], size=8, colour=MUTED)
    return sh


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # ---- header -----------------------------------------------------
    textbox(sl, 0.30, 0.22, 3.2, 0.3, [("GovStart Bridge  ·  SIH26136", 0)],
            size=12, colour=MUTED)
    textbox(sl, 3.55, 0.16, 9.4, 0.7, [("TECHNICAL APPROACH", 0)],
            size=30, align=PP_ALIGN.CENTER, colour=INK)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.30), Inches(0.82),
                                 Inches(13.03), Inches(0.82))
    ln.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); ln.line.width = Pt(1)

    # ---- left column ------------------------------------------------
    x, w = 0.30, 3.15
    y, gap = 0.98, 0.10
    bh = 1.06
    for title, tint, runs in BLOCKS:
        card = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(bh))
        card.fill.solid(); card.fill.fore_color.rgb = tint
        card.line.fill.background(); card.shadow.inherit = False
        card.text_frame.text = ""
        bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                  Inches(0.055), Inches(bh))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x59, 0x59, 0x59)
        bar.line.fill.background(); bar.shadow.inherit = False
        bar.text_frame.text = ""
        textbox(sl, x + 0.16, y + 0.07, w - 0.30, 0.22, [(title, 0)],
                size=11, colour=INK).text_frame.paragraphs[0].runs[0].font.bold = True
        textbox(sl, x + 0.16, y + 0.30, w - 0.30, bh - 0.36, runs,
                size=8.2, colour=MUTED, space=1.02)
        y += bh + gap

    # ---- the architecture -------------------------------------------
    fx, fw = 3.70, 9.33
    fy = 0.98
    fh = 1.18
    for title, nodes in STAGES:
        frame(sl, fx, fy, fw, fh, title)
        n = len(nodes)
        inner = fw - 0.34
        gapx = 0.30
        nw = (inner - gapx * (n - 1)) / n
        nh = 0.60
        ny = fy + 0.42
        nx = fx + 0.17
        for i, (label, fill) in enumerate(nodes):
            if fill is None:
                d = min(nw, 1.22)                      # wide enough not to
                node(sl, nx + (nw - d) / 2, ny - 0.10, d, nh + 0.20, label, None)
            else:
                node(sl, nx, ny, nw, nh, label, fill)
            if i < n - 1:
                arrow(sl, nx + nw + 0.04, ny + nh / 2, nx + nw + gapx - 0.04)
            nx += nw + gapx
        fy += fh + 0.10

    # ---- planned, not built -----------------------------------------
    py = fy + 0.02
    frame(sl, fx, py, fw, 0.52, "PLANNED INTEGRATIONS — NOT IN THE PROTOTYPE")
    n = len(PLANNED_ITEMS)
    inner = fw - 0.34; gapx = 0.22
    nw = (inner - gapx * (n - 1)) / n
    nx = fx + 0.17
    for item in PLANNED_ITEMS:
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(nx),
                                 Inches(py + 0.20), Inches(nw), Inches(0.26))
        sh.fill.solid(); sh.fill.fore_color.rgb = PLANNED
        sh.line.color.rgb = RGBColor(0xA6, 0xA6, 0xA6); sh.line.width = Pt(0.75)
        sh.line.dash_style = 4
        sh.shadow.inherit = False
        tf = sh.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.03)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = item
        r.font.name = "Calibri"; r.font.size = Pt(7.5); r.font.color.rgb = MUTED
        nx += nw + gapx

    # ---- key --------------------------------------------------------
    ky = py + 0.62
    legend = [("Platform module", PLATFORM), ("External system", EXTERNAL),
              ("Legal instrument", LEGAL), ("Governance checkpoint", CHECK),
              ("Planned", PLANNED)]
    lx = fx + 0.17
    for label, col in legend:
        sw = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lx), Inches(ky),
                                 Inches(0.20), Inches(0.13))
        sw.fill.solid(); sw.fill.fore_color.rgb = col
        sw.line.color.rgb = LINE; sw.line.width = Pt(0.5); sw.shadow.inherit = False
        sw.text_frame.text = ""
        textbox(sl, lx + 0.26, ky - 0.02, 1.7, 0.2, [(label, 0)], size=7.5, colour=MUTED)
        lx += 1.86

    # ---- foot band ---------------------------------------------------
    band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.30), Inches(6.92),
                               Inches(12.73), Inches(0.42))
    band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    band.line.fill.background(); band.shadow.inherit = False
    tf = band.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.14); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("No amendment to the General Financial Rules. One state Government "
              "Resolution unlocks Tier 1 — and Tiers 2 and 3 work without even that.")
    r.font.name = "Calibri"; r.font.size = Pt(11.5); r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    prs.save(OUT)
    print("wrote", OUT)


build()
