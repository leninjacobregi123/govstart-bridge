# -*- coding: utf-8 -*-
"""Build the 'Impact and benefits' slide.

Organised by who benefits rather than by adjective, because "impact" on its
own persuades nobody: the question a judge asks is who is better off and in
what way. The three bands underneath group the same benefits the way SIH
usually asks for them - social, economic, governance.

The claim is the case for adoption. What exists today is a working
prototype of the mechanism, not measured field results, and the slide says
so rather than implying otherwise.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from slide_technical import textbox, INK, MUTED, ACCENT

OUT = "/home/lenin/Apps Developed/SIH 26136/slide-assets/impact-and-benefits.pptx"
NAVY = RGBColor(0x1F, 0x3B, 0x63)

LEAD = [
 ("By letting a department buy the ", 0), ("evidence", 1),
 (" about an innovation before it buys the product, problems citizens actually "
  "feel turn from ", 0), ("unprocurable into solvable", 1),
 (" — and because nothing here amends a rule, the gain is available this "
  "financial year, in any state that wants it.", 0),
]

WHO = [
 ("Citizens", RGBColor(0xE7,0xF0,0xFA),
  [("The problems that currently get written off. OPD waiting time, "
    "non-revenue water, grievance backlogs and last-mile delivery in "
    "Gadchiroli and Nandurbar become ", 0), ("buyable", 1),
   (", because a department can at last contract for the thing that fixes "
    "them.", 0)]),
 ("Departments and the exchequer", RGBColor(0xE9,0xF3,0xEC),
  [("Money follows proof. The Evidence Contract is ", 0),
   ("milestone-paid and capped", 1), (", so a trial that fails costs a "
    "fraction of a rollout that fails — and criteria sealed before the "
    "result make the award ", 0), ("defensible in audit", 1), (".", 0)]),
 ("Startups and small firms", RGBColor(0xFD,0xF6,0xE3),
  [("A first-time supplier with ", 0), ("no turnover and no prior contract", 1),
   (" can win public revenue on merit. Milestone payment and a seven-day "
    "grievance clock meet the cash-flow problem that kills small firms.", 0)]),
 ("The state — and every other state", RGBColor(0xF3,0xEA,0xF9),
  [("No amendment to the General Financial Rules. One Government Resolution "
    "unlocks Tier 1; Tiers 2 and 3 need nothing. ", 0),
   ("What Maharashtra proves, another state can copy", 1),
   (" without legislating.", 0)]),
]

BANDS = [
 ("SOCIAL",
  [("Services that reach the people who need them, including the districts "
    "usually reached last. The platform itself is ", 0),
   ("WCAG AA and bilingual", 1), (", so the public it serves can use it.", 0)]),
 ("ECONOMIC",
  [("A wider supplier base, skilled local jobs, and public spend that buys ", 0),
   ("outcomes instead of promises", 1), (" — with the downside capped at the "
    "cost of the evidence, not the rollout.", 0)]),
 ("GOVERNANCE",
  [("Decisions that survive scrutiny: criteria sealed before results, a "
    "validator named in advance, and a route finder that ", 0),
   ("refuses rather than improvises", 1), (" when no lawful route exists.", 0)]),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    textbox(sl, 0.30, 0.22, 3.6, 0.3, [("GovStart Bridge  ·  SIH26136", 0)],
            size=12, colour=MUTED)
    textbox(sl, 3.55, 0.16, 9.4, 0.7, [("IMPACT AND BENEFITS", 0)],
            size=30, align=PP_ALIGN.CENTER, colour=INK)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.30), Inches(0.82),
                                 Inches(13.03), Inches(0.82))
    ln.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); ln.line.width = Pt(1)

    textbox(sl, 0.30, 0.92, 12.73, 0.46, LEAD, size=11.5, colour=MUTED,
            align=PP_ALIGN.CENTER, space=1.10)

    # ---- who is better off, and how ---------------------------------
    y, ch = 1.58, 2.18
    n = len(WHO); gap = 0.20
    cw = (12.73 - gap * (n - 1)) / n
    x = 0.30
    for title, tint, runs in WHO:
        card = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(cw), Inches(ch))
        card.fill.solid(); card.fill.fore_color.rgb = tint
        card.line.fill.background(); card.shadow.inherit = False
        card.text_frame.text = ""
        rule = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(cw), Inches(0.05))
        rule.fill.solid(); rule.fill.fore_color.rgb = NAVY
        rule.line.fill.background(); rule.shadow.inherit = False
        rule.text_frame.text = ""
        t = textbox(sl, x + 0.18, y + 0.20, cw - 0.36, 0.50, [(title, 0)],
                    size=12.5, colour=INK)
        t.text_frame.paragraphs[0].runs[0].font.bold = True
        textbox(sl, x + 0.18, y + 0.74, cw - 0.36, ch - 0.90, runs,
                size=10, colour=MUTED, space=1.10)
        x += cw + gap

    # ---- the same benefits, grouped the way SIH asks -----------------
    y2, bh = 4.06, 1.78
    n = len(BANDS); gap = 0.22
    bw = (12.73 - gap * (n - 1)) / n
    x = 0.30
    for title, runs in BANDS:
        box = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y2),
                                  Inches(bw), Inches(bh))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xF9)
        box.line.color.rgb = RGBColor(0xD6, 0xDC, 0xE2); box.line.width = Pt(0.75)
        box.shadow.inherit = False; box.text_frame.text = ""
        t = textbox(sl, x + 0.22, y2 + 0.20, bw - 0.44, 0.26, [(title, 0)],
                    size=10.5, colour=NAVY)
        t.text_frame.paragraphs[0].runs[0].font.bold = True
        textbox(sl, x + 0.22, y2 + 0.56, bw - 0.44, bh - 0.76, runs,
                size=10, colour=MUTED, space=1.10)
        x += bw + gap

    foot = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.30), Inches(6.22),
                               Inches(12.73), Inches(0.42))
    foot.fill.solid(); foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background(); foot.shadow.inherit = False
    tf = foot.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("Problems that were unprocurable become solvable — this financial "
              "year, without changing a single rule.")
    r.font.name = "Calibri"; r.font.size = Pt(11.5); r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    textbox(sl, 0.30, 6.80, 12.73, 0.24,
            [("This is the case for adoption. What exists today is a working "
              "prototype of the mechanism — not measured field results.", 0)],
            size=8, colour=RGBColor(0x8A, 0x8A, 0x8A), align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
