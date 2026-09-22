# -*- coding: utf-8 -*-
"""Build the 'Key highlights' slide: five problems, five answers.

Each row is one thing that stops a department buying from a startup today,
and the one thing this mechanism does about it. Every answer on the right is
something the prototype actually does; nothing here is aspirational.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from slide_technical import textbox, arrow, INK, MUTED, ACCENT, LINE

OUT = "/home/lenin/Apps Developed/SIH 26136/slide-assets/key-highlights.pptx"

PROB_TINT = RGBColor(0xFC, 0xEC, 0xEA)
SOLV_TINT = RGBColor(0xE9, 0xF3, 0xEC)
NAVY      = RGBColor(0x1F, 0x3B, 0x63)

ROWS = [
 ("Cannot buy what cannot be specified",
  [("Procurement rules assume you can describe the thing before you buy it. "
    "An innovation is the one purchase where you cannot.", 0)],
  "Buy the evidence first, the product second",
  [("An ", 0), ("Evidence Contract", 1), (" buys a defined question, test and report. "
    "The department buys the product afterwards — once that evidence has made it "
    "specifiable.", 0)]),

 ("The first-time supplier is filtered out",
  [("Turnover and prior-experience conditions disqualify the startup that can "
    "actually solve it, before anyone reads the idea.", 0)],
  "Eligibility from capped risk, not turnover",
  [("Five axes set a measured risk cap, and eligibility follows the cap. That is "
    "the exact condition ", 0), ("Rule 173(i)", 1),
   (" attaches to relaxing the filter.", 0)]),

 ("The goalposts move after the results",
  [("Criteria get “clarified” once everyone can see who would win. No startup "
    "can invest against that, and no officer can defend it.", 0)],
  "Criteria sealed before any solution is seen",
  [("The criteria are hashed with ", 0), ("SHA-256", 1), (" and published with the "
    "challenge. Edit one target afterwards and the seal ", 0),
   ("visibly breaks in front of everyone", 1), (".", 0)]),

 ("A pilot that works still dies there",
  [("Rule 166 blocks the direct award — winning a challenge is not a ground. "
    "Rule 157 makes pilot-then-rollout a piecemeal purchase.", 0)],
  "Three lawful routes, mapped in advance",
  [("PAC, limited tender to the winners, or GeM catalogue — chosen from facts "
    "already held. Where none of the three fits, the platform ", 0),
   ("says so instead of inventing one", 1), (".", 0)]),

 ("A fix that needs new law never arrives",
  [("Most proposals in this space require an amendment somebody has to pass "
    "first, so nothing changes this year or next.", 0)],
  "Nothing here amends the rules",
  [("No change to the General Financial Rules. ", 0),
   ("One state Government Resolution", 1), (" unlocks Tier 1 — and Tiers 2 and 3 "
    "need nothing at all, so a state can start next quarter.", 0)]),
]


def cell(sl, x, y, w, h, tint, head, runs, head_colour):
    box = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                              Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = tint
    box.line.fill.background(); box.shadow.inherit = False
    box.text_frame.text = ""
    t = textbox(sl, x + 0.18, y + 0.11, w - 0.36, 0.24, [(head, 0)],
                size=11.5, colour=head_colour)
    t.text_frame.paragraphs[0].runs[0].font.bold = True
    textbox(sl, x + 0.18, y + 0.37, w - 0.36, h - 0.46, runs,
            size=9.2, colour=MUTED, space=1.05)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    textbox(sl, 0.30, 0.22, 3.6, 0.3, [("GovStart Bridge  ·  SIH26136", 0)],
            size=12, colour=MUTED)
    textbox(sl, 3.55, 0.16, 9.4, 0.7, [("KEY HIGHLIGHTS", 0)],
            size=30, align=PP_ALIGN.CENTER, colour=INK)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.30), Inches(0.82),
                                 Inches(13.03), Inches(0.82))
    ln.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); ln.line.width = Pt(1)

    # column headings
    t = textbox(sl, 0.72, 0.90, 5.2, 0.24, [("THE PROBLEM TODAY", 0)],
                size=9.5, colour=RGBColor(0xA8, 0x3A, 0x2E))
    t.text_frame.paragraphs[0].runs[0].font.bold = True
    t = textbox(sl, 7.00, 0.90, 6.0, 0.24, [("WHAT GOVSTART BRIDGE DOES", 0)],
                size=9.5, colour=RGBColor(0x1E, 0x6B, 0x45))
    t.text_frame.paragraphs[0].runs[0].font.bold = True

    y = 1.18
    rh, gap = 1.06, 0.10
    px, pw = 0.72, 5.20
    sx, sw = 7.00, 6.03
    for i, (ph, pr, sh_, sr) in enumerate(ROWS):
        chip = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.30), Inches(y + 0.33),
                                   Inches(0.32), Inches(0.32))
        chip.fill.solid(); chip.fill.fore_color.rgb = NAVY
        chip.line.fill.background(); chip.shadow.inherit = False
        ctf = chip.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run(); cr.text = str(i + 1)
        cr.font.size = Pt(11.5); cr.font.bold = True; cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        cell(sl, px, y, pw, rh, PROB_TINT, ph, pr, RGBColor(0xA8, 0x3A, 0x2E))
        arrow(sl, px + pw + 0.14, y + rh / 2, sx - 0.14)
        cell(sl, sx, y, sw, rh, SOLV_TINT, sh_, sr, RGBColor(0x1E, 0x6B, 0x45))
        y += rh + gap

    foot = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.30), Inches(6.92),
                               Inches(12.73), Inches(0.42))
    foot.fill.solid(); foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background(); foot.shadow.inherit = False
    tf = foot.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("Five blockers, five answers — and not one of them needs a change "
              "to the General Financial Rules.")
    r.font.name = "Calibri"; r.font.size = Pt(11.5); r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
